#!/usr/bin/env python3
"""
metadata.py - collect and present metadata for video and image files.

Replaces the earlier meta.py, metas.py, metas2.py and metax.py.
Run "metadata.py --help" for usage, "metadata.py --check" to verify prerequisites.
"""

import sys

if sys.version_info < (3, 8):
    sys.exit("metadata.py needs Python 3.8 or newer (found %s). Download: https://www.python.org/downloads/"
             % sys.version.split()[0])

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

__version__ = "1.1.0"
__date__ = "2026-10-26"
PROG = "metadata.py"

VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".m4v", ".webm", ".flv",
                    ".wmv", ".mpg", ".mpeg", ".ts", ".m2ts", ".3gp"}
IMAGE_FORMAT_FOR_EXT = {".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".gif": "GIF",
                        ".bmp": "BMP", ".tif": "TIFF", ".tiff": "TIFF", ".webp": "WEBP"}
IMAGE_EXTENSIONS = set(IMAGE_FORMAT_FOR_EXT)

# extension -> name ffprobe uses for the container, where they differ
CONTAINER_ALIAS = {"mkv": "matroska", "ts": "mpegts", "m2ts": "mpegts", "wmv": "asf",
                   "mpg": "mpeg", "m4v": "mov"}

FFPROBE_TIMEOUT = 60  # seconds per file

# (header, record key) - shared by the csv and xlsx writers
COLUMNS = [
    ("Folder", "folder"), ("File", "name"), ("Type", "type"),
    ("Format", "format"), ("Width", "width"), ("Height", "height"),
    ("Duration", "duration_text"), ("Video codec / colour mode", "codec"),
    ("Bitrate (kb/s)", "bitrate_kbps"), ("FPS", "fps"),
    ("Audio codec", "audio"), ("Size (bytes)", "bytes"),
    ("Modified", "modified"), ("Date taken", "taken"),
    ("Camera", "camera"), ("Note", "note"), ("Error", "error"),
]

# Which tools are usable; filled in by main() after the prerequisite check.
TOOLS = {"video": False, "image": False}


def warn(msg):
    print(msg, file=sys.stderr)


# --------------------------------------------------------------------------
# prerequisites
# --------------------------------------------------------------------------

def pip_hint(package):
    return f'"{sys.executable}" -m pip install {package}'


def ffmpeg_hint():
    if sys.platform.startswith("win"):
        return ("winget install Gyan.FFmpeg      (or: choco install ffmpeg / scoop install ffmpeg)\n"
                "      Then open a NEW terminal so the updated PATH is picked up.\n"
                "      Manual download: https://ffmpeg.org/download.html (add the 'bin' folder to PATH)")
    if sys.platform == "darwin":
        return "brew install ffmpeg"
    return ("sudo apt install ffmpeg   (Debian/Ubuntu)\n"
            "      sudo dnf install ffmpeg   (Fedora)\n"
            "      sudo pacman -S ffmpeg     (Arch)")


def module_version(name):
    """Return the version string of an importable module, or None if unusable."""
    try:
        mod = __import__(name)
        return str(getattr(mod, "__version__", "unknown"))
    except Exception:
        return None


def ffprobe_version():
    exe = shutil.which("ffprobe")
    if not exe:
        return None
    try:
        out = subprocess.run([exe, "-version"], capture_output=True, text=True,
                             timeout=15, errors="replace").stdout
        return (out.splitlines() or ["ffprobe (unknown version)"])[0]
    except Exception:
        return None


def prerequisite_status():
    """name -> (version or None, what needs it, install hint)"""
    return {
        "ffprobe (FFmpeg)": (ffprobe_version(), "video files",
                             ffmpeg_hint()),
        "Pillow": (module_version("PIL"), "image files",
                   pip_hint("pillow")),
        "openpyxl": (module_version("openpyxl"), "Excel (xlsx) output",
                     pip_hint("openpyxl")),
    }


def run_check():
    """--check: print the state of every prerequisite. Returns exit code."""
    print(f"{PROG} {__version__} ({__date__})")
    print(f"Python {sys.version.split()[0]} at {sys.executable}\n")
    missing = 0
    for name, (version, needed_for, hint) in prerequisite_status().items():
        if version:
            print(f"  [ OK ]  {name}: {version}")
        else:
            missing += 1
            print(f"  [MISSING]  {name} - needed for {needed_for}\n      Install: {hint}")
    print("\nAll prerequisites are installed." if not missing else
          f"\n{missing} prerequisite(s) missing. {PROG} only needs the ones for the files/outputs you use.")
    return 1 if missing else 0


def require(need_video, need_image, need_xlsx):
    """Exit with install instructions if something needed is missing; set TOOLS."""
    status = prerequisite_status()
    TOOLS["video"] = status["ffprobe (FFmpeg)"][0] is not None
    TOOLS["image"] = status["Pillow"][0] is not None
    needed = []
    if need_video:
        needed.append("ffprobe (FFmpeg)")
    if need_image:
        needed.append("Pillow")
    if need_xlsx:
        needed.append("openpyxl")
    missing = [n for n in needed if status[n][0] is None]
    if not missing:
        return
    lines = ["Missing prerequisites:"]
    for n in missing:
        lines.append(f"  - {n} (needed for {status[n][1]})\n      Install: {status[n][2]}")
    lines.append(f"\nRun '{PROG} --check' to re-test after installing.")
    if "Pillow" in missing or "ffprobe (FFmpeg)" in missing:
        lines.append("Tip: to skip a kind of file, use --type video or --type image.")
    sys.exit("\n".join(lines))


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def human_size(n):
    n = float(n)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024.0:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024.0
    return f"{n:.1f} PB"


def format_duration(seconds):
    s = int(round(seconds))
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


def parse_fraction(text):
    try:
        num, den = text.split("/")
        return float(num) / float(den) if float(den) else None
    except (ValueError, AttributeError):
        return None


def to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def to_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def friendly_error(e):
    """Turn an exception into a short message a user can act on."""
    if isinstance(e, PermissionError):
        return "Access denied or file is locked by another program"
    if isinstance(e, FileNotFoundError):
        return "File not found (moved or deleted during the scan)"
    if isinstance(e, subprocess.TimeoutExpired):
        return f"Timed out after {FFPROBE_TIMEOUT} s while reading the file"
    if isinstance(e, json.JSONDecodeError):
        return "ffprobe returned unreadable output"
    if type(e).__name__ == "UnidentifiedImageError":
        return "Not a recognized image (corrupt or unsupported)"
    if "truncated" in str(e).lower():
        return "File is truncated or damaged"
    return str(e).strip() or type(e).__name__


# --------------------------------------------------------------------------
# metadata extraction
# --------------------------------------------------------------------------

def sniff_is_image(head):
    """True if the first bytes of a file are the signature of a supported image."""
    return (head.startswith((b"\xff\xd8\xff", b"\x89PNG\r\n\x1a\n", b"GIF87a", b"GIF89a",
                             b"II*\x00", b"MM\x00*"))
            or (head[:2] == b"BM" and head[6:10] == b"\0\0\0\0")
            or (head[:4] == b"RIFF" and head[8:12] == b"WEBP"))


def probe_video(path):
    """Return a dict of metadata; raise on failure."""
    cmd = ["ffprobe", "-v", "error", "-show_streams", "-show_format", "-of", "json", str(path)]
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                         errors="replace", timeout=FFPROBE_TIMEOUT)
    if res.returncode != 0:
        msg = (res.stderr.strip().splitlines() or ["ffprobe could not read the file"])[-1]
        raise ValueError(msg.replace(f"{path}: ", ""))
    data = json.loads(res.stdout or "{}")
    streams = data.get("streams", [])
    fmt = data.get("format", {})

    video = next((s for s in streams if s.get("codec_type") == "video"
                  and not s.get("disposition", {}).get("attached_pic")), None)
    audio = next((s for s in streams if s.get("codec_type") == "audio"), None)
    if video is None:
        raise ValueError("No video stream in file")

    ext = path.suffix[1:].lower()
    tokens = (fmt.get("format_name") or "").split(",")
    if any(t == "image2" or t.endswith("_pipe") for t in tokens):
        # ffprobe guesses "still image" from the extension even for damaged data
        raise ValueError("Not a valid video (damaged or not a video file)")
    ext_ok = CONTAINER_ALIAS.get(ext, ext) in tokens
    out = {
        "format": ext.upper() if ext_ok else (tokens[0] or "unknown").upper(),
        "_ext_ok": ext_ok, "_container": ",".join(tokens),
        "width": video.get("width"), "height": video.get("height"),
        "codec": video.get("codec_name"),
        "audio": audio.get("codec_name") if audio else "none",
    }
    duration = to_float(fmt.get("duration")) or to_float(video.get("duration"))
    if duration is not None:
        out["duration"] = duration
        out["duration_text"] = format_duration(duration)
    bitrate = to_int(fmt.get("bit_rate")) or to_int(video.get("bit_rate"))
    if bitrate:
        out["bitrate_kbps"] = round(bitrate / 1000)
    fps = parse_fraction(video.get("avg_frame_rate")) or parse_fraction(video.get("r_frame_rate"))
    if fps:
        out["fps"] = round(fps, 2)
    return out


def probe_image(path):
    """Return a dict of metadata; raise on failure."""
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = None  # only the header is read; don't trip the bomb guard
    with Image.open(path) as img:
        out = {"format": img.format, "width": img.size[0], "height": img.size[1],
               "codec": img.mode,
               "_ext_ok": IMAGE_FORMAT_FOR_EXT.get(path.suffix.lower()) == img.format}
        try:
            exif = img.getexif()
            taken = exif.get_ifd(0x8769).get(36867) or exif.get(306)  # DateTimeOriginal / DateTime
            if taken:
                out["taken"] = str(taken).strip()
            camera = " ".join(str(exif[t]).strip() for t in (271, 272) if exif.get(t))
            if camera:
                out["camera"] = camera
        except Exception:
            pass  # missing or damaged EXIF is not an error
    return out


PROBES = {"video": probe_video, "image": probe_image}


def analyze(path, root, labeled):
    """Analyze one file. Never raises: problems are recorded in rec["error"]."""
    rel = path.relative_to(root)
    rec = {"path": path, "relpath": str(rel), "folder": str(rel.parent),
           "name": path.name, "type": labeled, "bytes": 0}
    try:
        st = path.stat()
        rec["bytes"] = st.st_size
        rec["modified"] = datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
        if st.st_size == 0:
            rec["error"] = "Empty file (0 bytes)"
            return rec
        with open(path, "rb") as f:  # also catches locked / unreadable files early
            head = f.read(32)
    except Exception as e:
        rec["error"] = friendly_error(e)
        return rec

    other = "image" if labeled == "video" else "video"
    order = ["image", "video"] if sniff_is_image(head) else [labeled, other]
    first_error = None
    for kind in order:
        if not TOOLS[kind]:
            continue
        try:
            fields = PROBES[kind](path)
        except Exception as e:
            first_error = first_error or friendly_error(e)
            continue
        ext_ok = fields.pop("_ext_ok")
        container = fields.pop("_container", None)
        rec.update(fields)
        rec["type"] = kind
        ext = path.suffix.lower()
        what = f"{fields['format']} {kind}"
        if kind != labeled:
            rec["note"] = f"Mislabeled: extension {ext} suggests {labeled}, but the content is {what}"
        elif not ext_ok:
            rec["note"] = f"Extension {ext} does not match the content ({container or fields['format']})"
        return rec
    rec["error"] = first_error or "Unsupported or corrupt file"
    return rec


# --------------------------------------------------------------------------
# file discovery and ordering
# --------------------------------------------------------------------------

def find_files(root, recursive, types):
    exts = {}
    if types in ("video", "all"):
        exts.update({e: "video" for e in VIDEO_EXTENSIONS})
    if types in ("image", "all"):
        exts.update({e: "image" for e in IMAGE_EXTENSIONS})

    def skipped(err):
        warn(f"Warning: cannot read folder {err.filename}: {friendly_error(err)}")

    candidates = []
    try:
        if recursive:
            for d, _, files in os.walk(root, onerror=skipped):
                candidates.extend(Path(d) / f for f in files)
        else:
            candidates = list(root.iterdir())
    except OSError as e:
        sys.exit(f"Cannot read folder {root}: {friendly_error(e)}")

    found = []
    for p in candidates:
        kind = exts.get(p.suffix.lower())
        if kind:
            try:
                if p.is_file():
                    found.append((p, kind))
            except OSError:
                pass
    return found


def sort_records(records, mode):
    if mode == "folder":
        return sorted(records, key=lambda r: (r["folder"].lower(), r["name"].lower()))
    return sorted(records, key=lambda r: (r["name"].lower(), r["folder"].lower()))


# --------------------------------------------------------------------------
# rename
# --------------------------------------------------------------------------

def unique_path(path):
    if not path.exists():
        return path
    for i in range(1, 1000):
        cand = path.with_name(f"{path.stem}_{i:02d}{path.suffix}")
        if not cand.exists():
            return cand
    return path


def rename_record(rec, dry_run):
    if "error" in rec or not rec.get("width"):
        return
    suffix = f" ({rec['width']}x{rec['height']}"
    if rec["type"] == "video" and rec.get("duration_text"):
        suffix += " - " + rec["duration_text"].replace(":", "-")
    suffix += ")"
    old = rec["path"]
    if old.stem.endswith(suffix):
        return  # already renamed
    new = unique_path(old.with_name(old.stem + suffix + old.suffix))
    warn(f"{'Would rename' if dry_run else 'Renaming'}: {old.name} -> {new.name}")
    if dry_run:
        return
    try:
        old.rename(new)
    except OSError as e:
        warn(f"  not renamed: {friendly_error(e)}")
        rec["note"] = (rec.get("note", "") + " " if rec.get("note") else "") + \
            f"Rename failed: {friendly_error(e)}"
        return
    rec["path"] = new
    rec["name"] = new.name
    rec["relpath"] = str(Path(rec["folder"]) / new.name)


# --------------------------------------------------------------------------
# output
# --------------------------------------------------------------------------

def record_text(rec, show_path):
    lines = [f"File: {rec['relpath'] if show_path else rec['name']}"]
    if "error" in rec:
        lines.append(f"  Error: {rec['error']}")
    else:
        lines.append(f"  Type: {rec['type']} ({rec.get('format') or 'N/A'})")
        lines.append(f"  Resolution: {rec.get('width', 'N/A')}x{rec.get('height', 'N/A')}")
        if rec["type"] == "video":
            lines.append(f"  Duration: {rec.get('duration_text', 'N/A')}")
            lines.append(f"  Video codec: {rec.get('codec', 'N/A')}")
            lines.append(f"  Audio codec: {rec.get('audio', 'N/A')}")
            if "bitrate_kbps" in rec:
                lines.append(f"  Bitrate: {rec['bitrate_kbps']} kb/s")
            if "fps" in rec:
                lines.append(f"  Frame rate: {rec['fps']} fps")
        else:
            lines.append(f"  Colour mode: {rec.get('codec', 'N/A')}")
            if "taken" in rec:
                lines.append(f"  Date taken: {rec['taken']}")
            if "camera" in rec:
                lines.append(f"  Camera: {rec['camera']}")
    if "note" in rec:
        lines.append(f"  Note: {rec['note']}")
    lines.append(f"  Size: {human_size(rec['bytes'])} ({rec['bytes']} bytes)")
    if "modified" in rec:
        lines.append(f"  Modified: {rec['modified']}")
    return "\n".join(lines) + "\n"


def summary_text(label, recs):
    errors = sum(1 for r in recs if "error" in r)
    notes = sum(1 for r in recs if "note" in r)
    videos = sum(1 for r in recs if r["type"] == "video")
    text = (f"{label}: {len(recs)} file{'s' * (len(recs) != 1)} "
            f"({videos} video, {len(recs) - videos} image), "
            f"{human_size(sum(r['bytes'] for r in recs))}")
    if errors:
        text += f", {errors} with errors"
    if notes:
        text += f", {notes} with notes"
    return text


def render_text(records, grouped, root, blank=True, title=""):
    out = [f"Metadata for {root}{title}",
           f"Generated by {PROG} {__version__} on {datetime.now():%Y-%m-%d %H:%M:%S}", ""]
    if grouped:
        folders = {}
        for r in records:
            folders.setdefault(r["folder"], []).append(r)
        for folder, recs in folders.items():
            out.append(f"Folder: {folder}")
            out.append("")
            out.extend(record_text(r, False) for r in recs)
            out.append(summary_text("Folder summary", recs))
            if blank:
                out.append("")
    else:
        show_path = any(r["folder"] != "." for r in records)
        out.extend(record_text(r, show_path) for r in records)
    out.append(summary_text("Grand summary", records))
    return "\n".join(out) + "\n"


def row_values(rec):
    return [rec.get(key) for _, key in COLUMNS]


def write_csv(path, records, delimiter, blank_between=False):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow([h for h, _ in COLUMNS])
        prev = None
        for r in records:
            if blank_between and prev is not None and r["folder"] != prev:
                w.writerow([])
            prev = r["folder"]
            w.writerow(["" if v is None else v for v in row_values(r)])


def write_xlsx(path, records, blank_between=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.properties.title = "Media metadata"
    wb.properties.creator = f"{PROG} {__version__}"
    ws = wb.active
    ws.title = "Metadata"
    bold = Font(bold=True)
    ws.append([h for h, _ in COLUMNS])
    for c in ws[1]:
        c.font = bold
    prev = None
    for r in records:
        if blank_between and prev is not None and r["folder"] != prev:
            ws.append([])
        prev = r["folder"]
        ws.append(row_values(r))
    size_col = [k for _, k in COLUMNS].index("bytes") + 1
    for (cell,) in ws.iter_rows(min_row=2, min_col=size_col, max_col=size_col):
        cell.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (header, key) in enumerate(COLUMNS, start=1):
        widest = max([len(header)] + [len(str(r.get(key) or "")) for r in records])
        ws.column_dimensions[get_column_letter(i)].width = min(widest + 2, 60)

    # per-folder summary sheet
    sm = wb.create_sheet("Summary")
    sm.append(["Folder", "Files", "Videos", "Images", "Errors", "Size (bytes)"])
    for c in sm[1]:
        c.font = bold
    folders = {}
    for r in records:
        folders.setdefault(r["folder"], []).append(r)
    for folder in sorted(folders, key=str.lower):
        recs = folders[folder]
        videos = sum(1 for r in recs if r["type"] == "video")
        sm.append([folder, len(recs), videos, len(recs) - videos,
                   sum(1 for r in recs if "error" in r),
                   sum(r["bytes"] for r in recs)])
    last = sm.max_row
    sm.append(["Total"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last})"
                           for c in range(2, 7)])
    for c in sm[sm.max_row]:
        c.font = bold
    for row in sm.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = "#,##0"
    sm.column_dimensions["A"].width = 50
    for col in "BCDEF":
        sm.column_dimensions[col].width = 14
    wb.save(path)


def output_path(name, ext, overwrite):
    path = Path(f"{name}.{ext}")
    if overwrite or not path.exists():
        return path
    return unique_path(path)


def save(path, writer):
    """Run writer(path). Returns True on success. A locked target gets a new name."""
    try:
        try:
            writer(path)
        except PermissionError:
            alt = unique_path(path)
            if alt == path:
                raise
            warn(f"'{path}' is locked (open in another program?), writing '{alt}' instead.")
            writer(alt)
            path = alt
    except Exception as e:
        reason = ("the output folder does not exist" if isinstance(e, FileNotFoundError)
                  else friendly_error(e))
        warn(f"Error: could not write '{path}': {reason}")
        return False
    warn(f"Wrote {path}")
    return True


# --------------------------------------------------------------------------
# command line
# --------------------------------------------------------------------------

HELP_EPILOG = f"""\
Switches are case-insensitive (-R = -r, --Recursive = --recursive) and every
option has a unix-style long form. A folder name may be given before or after
the options.

output formats (-o):
  screen   print to the terminal (default)
  txt      metadata.txt        csv   metadata.csv (Excel-friendly UTF-8)
  xlsx     metadata.xlsx       Several: -o txt,xlsx   or   -o txt -o xlsx

ordering (-s):
  name     by file name only (the folder is shown in front of the name)
  folder   by folder, then file name, with a summary per folder
           (default when -r is used; without -r there is only one folder)

kinds (-k):
  together  videos and images in one output (default)
  separate  one output per kind: two screen sections, and files named
            metadata_video.* and metadata_image.*

blank lines (-x):
  With ordering by folder, csv and xlsx get an empty row between folders so
  the groups stand out. Use -x to leave them out (one unbroken table, easier
  to sort/filter in Excel). In txt/screen it removes the gap after each
  folder summary.

mis-labeled and damaged files:
  The file *content* decides how a file is read, not its extension. A JPEG
  named clip.mp4 is reported as an image with a "Note". Locked, unreadable,
  empty or corrupt files are listed with an Error and never stop the run.

examples:
  {PROG}                          show the current folder on screen
  {PROG} -r -o xlsx               recursive, Excel sheet ordered by folder
  {PROG} D:\\Media -R -S name      recursive, ordered by file name only
  {PROG} --type image -o csv -d ";"   images only, CSV for European Excel
  {PROG} -n -p                    preview renaming (adds resolution/duration)
  {PROG} -r -k separate -o xlsx   separate Excel files for videos and images
  {PROG} -r -o csv -x             recursive CSV without blank rows between folders
  {PROG} --check                  verify prerequisites and show how to install

prerequisites:
  Python 3.8+, ffprobe from FFmpeg (video), Pillow (images), openpyxl (xlsx).
  Only the ones your files and chosen outputs need are required.

exit codes: 0 done, 1 error (missing prerequisite, nothing found, a file could
not be written), 2 bad command line, 130 interrupted.
"""


def parse_outputs(text):
    valid = ("screen", "txt", "csv", "xlsx")
    parts = [p.strip().lower() for p in text.split(",") if p.strip()]
    bad = [p for p in parts if p not in valid]
    if bad or not parts:
        raise argparse.ArgumentTypeError(
            f"invalid output {bad or text!r}; choose from {', '.join(valid)} (comma separated)")
    return parts


def normalize_argv(argv):
    """Make switches case-insensitive: -R -> -r, --Recursive -> --recursive."""
    out = []
    for i, a in enumerate(argv):
        if a == "--":
            out.extend(argv[i:])
            break
        if a.startswith("--"):
            name, eq, val = a.partition("=")
            out.append(name.lower() + eq + val)
        elif re.fullmatch(r"-[A-Za-z?]+", a):
            out.append(a.lower())
        else:
            out.append(a)
    return out


def build_parser():
    ap = argparse.ArgumentParser(
        prog=PROG, add_help=False,
        formatter_class=argparse.RawDescriptionHelpFormatter,
        description="Collect and present metadata from video and image files.",
        epilog=HELP_EPILOG)
    ap.add_argument("folder", nargs="?", default=".",
                    help="folder to analyze (default: current folder)")
    ap.add_argument("-r", "--recursive", action="store_true",
                    help="include subfolders")
    ap.add_argument("-o", "--output", action="append", type=parse_outputs, metavar="FORMAT",
                    help="screen (default), txt, csv, xlsx; comma separated or repeated")
    ap.add_argument("-s", "--sort", type=str.lower, choices=["name", "folder"],
                    help="name | folder (default: folder with -r, otherwise name)")
    ap.add_argument("-t", "--type", type=str.lower, choices=["video", "image", "all"],
                    default="all", help="video | image | all (default: all)")
    ap.add_argument("-k", "--kinds", type=str.lower, choices=["together", "separate"],
                    default="together",
                    help="together: videos and images in one output (default); "
                         "separate: one output per kind (metadata_video.*, metadata_image.*)")
    ap.add_argument("-x", "--no-blank-lines", action="store_true",
                    help="when ordered by folder, do not put a blank line between folders")
    ap.add_argument("-b", "--basename", default="metadata", metavar="NAME",
                    help="base name of output files (default: metadata)")
    ap.add_argument("-d", "--delimiter", default=",", metavar="CHAR",
                    help="CSV delimiter; ';' suits European Excel (default: ,)")
    ap.add_argument("-w", "--overwrite", action="store_true",
                    help="overwrite existing output files instead of adding _01, _02, ...")
    ap.add_argument("-n", "--rename", action="store_true",
                    help="rename files to add resolution (and duration) to the name")
    ap.add_argument("-p", "--dry-run", action="store_true",
                    help="with --rename: only show what would be renamed")
    ap.add_argument("-c", "--check", action="store_true",
                    help="check prerequisites, show how to install missing ones, and exit")
    ap.add_argument("-v", "--version", action="version",
                    version=f"{PROG} {__version__} ({__date__})")
    ap.add_argument("-h", "-?", "--help", action="help", help="show this help and exit")
    return ap


def run(args):
    if args.check:
        return run_check()

    root = Path(args.folder).resolve()
    if not root.exists():
        warn(f"Folder not found: {root}")
        return 1
    if not root.is_dir():
        warn(f"Not a folder: {root} (give a folder, not a file)")
        return 1
    if args.dry_run and not args.rename:
        warn("Note: --dry-run only matters together with --rename.")
    if len(args.delimiter) != 1:
        warn(f"CSV delimiter must be a single character, got {args.delimiter!r}.")
        return 2

    outputs = list(dict.fromkeys(f for group in (args.output or [["screen"]]) for f in group))
    mode = args.sort or ("folder" if args.recursive else "name")
    if mode == "folder" and not args.recursive and args.sort:
        warn("Note: --sort folder has no effect without --recursive.")

    # fail early, before the slow part, if Excel output cannot work
    files = find_files(root, args.recursive, args.type)
    if not files:
        warn("No video or image files found" + ("" if args.recursive else
             " (add -r to include subfolders)") + ".")
        return 1
    kinds = {k for _, k in files}
    require("video" in kinds, "image" in kinds, "xlsx" in outputs)

    show_progress = sys.stderr.isatty()
    warn(f"Analyzing {len(files)} file{'s' * (len(files) != 1)} in {root} ...")
    records = []
    for i, (p, k) in enumerate(files, 1):
        if show_progress:
            print(f"\r  {i}/{len(files)}", end="", file=sys.stderr, flush=True)
        records.append(analyze(p, root, k))
    if show_progress:
        print(file=sys.stderr)

    if args.rename:
        for rec in records:
            rename_record(rec, args.dry_run)

    records = sort_records(records, mode)
    grouped = args.recursive and mode == "folder"
    status = 0
    blank = not args.no_blank_lines
    if args.kinds == "separate":
        parts = [(k, f"_{k}", f" - {k}s") for k in ("video", "image")
                 if any(r["type"] == k for r in records)]
        sets = [([r for r in records if r["type"] == k], suffix, title)
                for k, suffix, title in parts]
    else:
        sets = [(records, "", "")]

    for recs, suffix, title in sets:
        text = None
        for fmt in outputs:
            base = args.basename + suffix
            if fmt == "screen":
                try:
                    print(render_text(recs, grouped, root, blank, title))
                except BrokenPipeError:
                    return status
                continue
            if fmt == "txt":
                text = text or render_text(recs, grouped, root, blank, title)
                ok = save(output_path(base, "txt", args.overwrite),
                          lambda p: p.write_text(text, encoding="utf-8"))
            elif fmt == "csv":
                ok = save(output_path(base, "csv", args.overwrite),
                          lambda p: write_csv(p, recs, args.delimiter, grouped and blank))
            else:
                ok = save(output_path(base, "xlsx", args.overwrite),
                          lambda p: write_xlsx(p, recs, grouped and blank))
            if not ok:
                status = 1

    problems = [r for r in records if "error" in r]
    if problems:
        warn(f"\n{len(problems)} file{'s' * (len(problems) != 1)} could not be read:")
        for r in problems[:20]:
            warn(f"  {r['relpath']}: {r['error']}")
        if len(problems) > 20:
            warn(f"  ... and {len(problems) - 20} more (see the output)")
    return status


def main():
    for stream in (sys.stdout, sys.stderr):  # odd file names must not crash printing
        try:
            stream.reconfigure(errors="replace")
        except AttributeError:
            pass
    args = build_parser().parse_args(normalize_argv(sys.argv[1:]))
    try:
        return run(args)
    except KeyboardInterrupt:
        warn("\nInterrupted.")
        return 130


if __name__ == "__main__":
    sys.exit(main())
